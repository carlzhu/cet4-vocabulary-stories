from __future__ import annotations

import csv
import hashlib
import json
import re
import unicodedata
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

SUPPORTED_SUFFIXES = {".csv", ".tsv", ".txt", ".json", ".xlsx"}
GENERATED_NAMES = {
    "curriculum_plan.csv",
    "manifest.json",
    "review_schedule.csv",
    "vocabulary_allocation.csv",
    "vocabulary_coverage.csv",
    "vocabulary_master.csv",
    "vocabulary_source.csv",
}

ALIASES = {
    "lemma": {
        "lemma", "word", "term", "vocabulary", "headword", "单词", "词汇", "词条",
    },
    "part_of_speech": {
        "part_of_speech", "part of speech", "pos", "词性",
    },
    "chinese_meaning": {
        "chinese_meaning", "chinese meaning", "meaning", "definition", "translation",
        "释义", "中文释义", "中文", "意思",
    },
    "word_family": {
        "word_family", "word family", "family", "词族", "同根词",
    },
}

SOURCE_FIELDS = [
    "source_file", "source_sha256", "source_row", "raw_record_json", "raw_lemma",
    "raw_part_of_speech", "raw_chinese_meaning", "raw_word_family",
]
MASTER_FIELDS = [
    "vocabulary_id", "lemma", "part_of_speech", "chinese_meaning", "word_family",
    "status", "assigned_article", "occurrence_count", "source_file", "source_row",
    "original_lemma", "duplicate_group", "needs_review",
]


@dataclass(frozen=True)
class InputRecord:
    row_number: int
    raw: dict[str, Any]


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFKC", str(value)).replace("\u00a0", " ")
    return re.sub(r"\s+", " ", text).strip()


def comparison_key(value: Any) -> str:
    return normalize_text(value).casefold()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def discover_candidates(root: Path) -> list[Path]:
    candidates: list[Path] = []
    for path in root.rglob("*"):
        if not path.is_file() or path.suffix.casefold() not in SUPPORTED_SUFFIXES:
            continue
        if path.name.casefold() in GENERATED_NAMES:
            continue
        if any(part.casefold() in {".git", ".kiro", ".venv", "build", "drafts"} for part in path.parts):
            continue
        candidates.append(path)
    return sorted(candidates)


def _read_delimited(path: Path, delimiter: str | None = None) -> list[InputRecord]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        sample = handle.read(8192)
        handle.seek(0)
        if delimiter is None:
            try:
                delimiter = csv.Sniffer().sniff(sample, delimiters=",\t;|").delimiter
            except csv.Error:
                delimiter = ","
        reader = csv.DictReader(handle, delimiter=delimiter)
        if not reader.fieldnames:
            return []
        return [InputRecord(index, dict(row)) for index, row in enumerate(reader, start=2)]


def _read_text(path: Path) -> list[InputRecord]:
    lines = path.read_text(encoding="utf-8-sig").splitlines()
    meaningful = [(index, line.strip()) for index, line in enumerate(lines, start=1) if line.strip()]
    if not meaningful:
        return []
    if any("\t" in line for _, line in meaningful):
        return _read_delimited(path, delimiter="\t")
    return [InputRecord(index, {"lemma": line}) for index, line in meaningful]


def _read_json(path: Path) -> list[InputRecord]:
    payload = json.loads(path.read_text(encoding="utf-8-sig"))
    if isinstance(payload, dict):
        for key in ("words", "vocabulary", "entries", "items", "data"):
            if isinstance(payload.get(key), list):
                payload = payload[key]
                break
        else:
            payload = [payload]
    if not isinstance(payload, list):
        raise ValueError("JSON vocabulary source must be a list or contain a list field")
    result: list[InputRecord] = []
    for index, item in enumerate(payload, start=1):
        result.append(InputRecord(index, item if isinstance(item, dict) else {"lemma": item}))
    return result


def _read_xlsx(path: Path) -> list[InputRecord]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("XLSX import requires openpyxl; run `uv sync --extra dev`") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = [normalize_text(value) for value in next(rows, ())]
    result: list[InputRecord] = []
    for index, values in enumerate(rows, start=2):
        if not any(value is not None and normalize_text(value) for value in values):
            continue
        result.append(InputRecord(index, dict(zip(headers, values, strict=False))))
    workbook.close()
    return result


def read_records(path: Path) -> list[InputRecord]:
    suffix = path.suffix.casefold()
    if suffix == ".csv":
        return _read_delimited(path)
    if suffix == ".tsv":
        return _read_delimited(path, delimiter="\t")
    if suffix == ".txt":
        return _read_text(path)
    if suffix == ".json":
        return _read_json(path)
    if suffix == ".xlsx":
        return _read_xlsx(path)
    raise ValueError(f"Unsupported vocabulary source: {path.suffix}")


def map_columns(records: Iterable[InputRecord]) -> dict[str, str]:
    records = list(records)
    headers = list(records[0].raw) if records else []
    mapped: dict[str, str] = {}
    for canonical, aliases in ALIASES.items():
        for header in headers:
            if comparison_key(header) in aliases:
                mapped[canonical] = header
                break
    if "lemma" not in mapped and len(headers) == 1:
        mapped["lemma"] = headers[0]
    if "lemma" not in mapped:
        raise ValueError(
            "Could not identify the lemma column. Expected one of: "
            + ", ".join(sorted(ALIASES["lemma"]))
        )
    return mapped


def _stable_id(signature: tuple[str, str, str, str]) -> str:
    joined = "\x1f".join(signature).encode("utf-8")
    return "cet4-" + hashlib.sha1(joined).hexdigest()[:12]


def _write_csv(path: Path, fields: list[str], rows: Iterable[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def import_vocabulary(source_path: Path, project_root: Path) -> dict[str, Any]:
    source_path = source_path.resolve()
    project_root = project_root.resolve()
    records = read_records(source_path)
    if not records:
        raise ValueError("Vocabulary source contains no usable records")
    columns = map_columns(records)
    source_hash = sha256_file(source_path)

    source_rows: list[dict[str, Any]] = []
    grouped: dict[tuple[str, str, str, str], list[dict[str, Any]]] = defaultdict(list)
    skipped_rows: list[int] = []

    for record in records:
        extracted = {
            field: normalize_text(record.raw.get(column, ""))
            for field, column in columns.items()
        }
        raw_lemma = extracted.get("lemma", "")
        source_rows.append({
            "source_file": source_path.name,
            "source_sha256": source_hash,
            "source_row": record.row_number,
            "raw_record_json": json.dumps(record.raw, ensure_ascii=False, sort_keys=True, default=str),
            "raw_lemma": raw_lemma,
            "raw_part_of_speech": extracted.get("part_of_speech", ""),
            "raw_chinese_meaning": extracted.get("chinese_meaning", ""),
            "raw_word_family": extracted.get("word_family", ""),
        })
        if not raw_lemma:
            skipped_rows.append(record.row_number)
            continue
        signature = (
            comparison_key(raw_lemma),
            comparison_key(extracted.get("part_of_speech", "")),
            comparison_key(extracted.get("chinese_meaning", "")),
            comparison_key(extracted.get("word_family", "")),
        )
        grouped[signature].append({
            "record": record,
            "lemma": raw_lemma,
            "part_of_speech": extracted.get("part_of_speech", ""),
            "chinese_meaning": extracted.get("chinese_meaning", ""),
            "word_family": extracted.get("word_family", ""),
        })

    lemma_signatures: dict[str, list[tuple[str, str, str, str]]] = defaultdict(list)
    for signature in grouped:
        lemma_signatures[signature[0]].append(signature)

    duplicate_groups = {
        signature: f"dup-{index:04d}"
        for index, signature in enumerate(
            sorted((key for key, values in grouped.items() if len(values) > 1)), start=1
        )
    }
    master_rows: list[dict[str, Any]] = []
    ambiguous_lemmas: list[str] = []

    for signature in sorted(grouped):
        entries = grouped[signature]
        first = entries[0]
        same_lemma_variants = lemma_signatures[signature[0]]
        ambiguous = len(same_lemma_variants) > 1
        missing_core_metadata = not first["part_of_speech"] or not first["chinese_meaning"]
        if ambiguous:
            ambiguous_lemmas.append(first["lemma"])
        master_rows.append({
            "vocabulary_id": _stable_id(signature),
            "lemma": first["lemma"],
            "part_of_speech": first["part_of_speech"],
            "chinese_meaning": first["chinese_meaning"],
            "word_family": first["word_family"],
            "status": "needs_review" if ambiguous or missing_core_metadata else "unassigned",
            "assigned_article": "",
            "occurrence_count": 0,
            "source_file": source_path.name,
            "source_row": "|".join(str(item["record"].row_number) for item in entries),
            "original_lemma": first["lemma"],
            "duplicate_group": duplicate_groups.get(signature, ""),
            "needs_review": "yes" if ambiguous or missing_core_metadata else "no",
        })

    _write_csv(project_root / "vocabulary_source.csv", SOURCE_FIELDS, source_rows)
    _write_csv(project_root / "vocabulary_master.csv", MASTER_FIELDS, master_rows)

    audit = {
        "source_file": str(source_path),
        "source_sha256": source_hash,
        "raw_records": len(records),
        "unique_targets": len(master_rows),
        "exact_duplicate_rows": sum(len(rows) - 1 for rows in grouped.values()),
        "duplicate_groups": len(duplicate_groups),
        "ambiguous_lemmas": sorted(set(ambiguous_lemmas), key=str.casefold),
        "skipped_blank_lemma_rows": skipped_rows,
        "mapped_columns": columns,
        "needs_review": sum(row["needs_review"] == "yes" for row in master_rows),
    }
    write_audit(project_root / "vocabulary_audit.md", audit)
    return audit


def write_audit(path: Path, audit: dict[str, Any]) -> None:
    mapped = "\n".join(f"- `{key}` <- `{value}`" for key, value in audit["mapped_columns"].items())
    ambiguous = "\n".join(f"- {lemma}" for lemma in audit["ambiguous_lemmas"]) or "- None"
    skipped = ", ".join(map(str, audit["skipped_blank_lemma_rows"])) or "None"
    content = f"""# Vocabulary Data Audit

## Source

- File: `{audit['source_file']}`
- SHA-256: `{audit['source_sha256']}`
- Raw records: {audit['raw_records']}
- Unique normalized targets: {audit['unique_targets']}
- Exact duplicate rows removed from master: {audit['exact_duplicate_rows']}
- Duplicate groups: {audit['duplicate_groups']}
- Entries needing review: {audit['needs_review']}
- Blank-lemma source rows skipped from master: {skipped}

## Column mapping

{mapped}

## Same-lemma variants requiring review

{ambiguous}

## Rules applied

- Raw records remain preserved in `vocabulary_source.csv` with source row and file hash.
- Exact normalized duplicates are collapsed only in `vocabulary_master.csv`; source rows remain traceable.
- Different parts of speech, meanings, or word-family values remain distinct targets.
- Missing part of speech or Chinese meaning is left blank and marked `needs_review`; no data is invented.
- Word families and derived forms are never merged automatically.
"""
    path.write_text(content, encoding="utf-8")
